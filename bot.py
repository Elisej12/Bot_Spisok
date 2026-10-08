import sqlite3
import random
import io
import time
import html
import os
from datetime import datetime, timedelta

from dotenv import load_dotenv
import telebot
from telebot import types, apihelper
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ==================== 1. НАСТРОЙКИ И КОНСТАНТЫ ====================

# Загрузка .env строго из папки, где лежит сам скрипт
script_dir = os.path.dirname(os.path.abspath(__file__))
load_dotenv(os.path.join(script_dir, '.env'))

TOKEN = os.getenv('BOT_TOKEN')
if not TOKEN:
    raise ValueError("❌ Токен бота не найден! Проверьте файл .env")

# Парсинг списка админов из переменной окружения
admin_ids_raw = os.getenv('ADMIN_IDS', '')
ADMIN_IDS = [int(x.strip()) for x in admin_ids_raw.split(',') if x.strip().isdigit()]

DB_NAME = "schedule.db"

COLOR_BADGES = ["🟦", "🟩", "🟧", "🟪", "🟨", "🟥", "🟫", "⚪"]
DIVIDER_LINE = "➖➖➖➖➖➖➖➖➖➖➖➖"

bot = telebot.TeleBot(TOKEN)

# ==================== ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ ====================

def safe_send_message(chat_id, text, **kwargs):
    """
    Безопасная отправка сообщений с защитой от Telegram Rate Limits (лимит 30 соб/сек).
    """
    try:
        res = bot.send_message(chat_id, text, **kwargs)
        time.sleep(0.05)  # Небольшая задержка для защиты от спам-фильтра Telegram
        return res
    except apihelper.ApiTelegramException as e:
        if e.error_code == 429:  # Ошибка Too Many Requests
            retry_after = e.result_json.get('parameters', {}).get('retry_after', 5)
            print(f"⚠️ Превышен лимит сообщений. Ожидание {retry_after} сек...")
            time.sleep(retry_after)
            return safe_send_message(chat_id, text, **kwargs)
        print(f"Ошибка отправки пользователю {chat_id}: {e}")
        return None
    except Exception as e:
        print(f"Сбой отправки пользователю {chat_id}: {e}")
        return None

def safe_html(text):
    """Безопасное экранирование спецсимволов для HTML-разметки Telegram"""
    if not text:
        return ""
    return html.escape(str(text))

def get_next_week_days():
    """Формирует список дней с датами следующей недели"""
    today = datetime.now()
    days_ahead = 0 - today.weekday()
    if days_ahead <= 0:
        days_ahead += 7
    next_monday = today + timedelta(days=days_ahead)
    
    days_names = ["Понедельник", "Вторник", "Среда", "Четверг", "Пятница", "Суббота", "Воскресенье"]
    week_days = []
    for i in range(7):
        day_date = next_monday + timedelta(days=i)
        date_str = day_date.strftime("%d.%m")
        week_days.append(f"{days_names[i]} ({date_str})")
    return week_days

def get_user_role_text(user_id):
    if user_id in ADMIN_IDS:
        return "👑 <b>Администратор</b>"
    return "👤 <b>Сотрудник</b>"

def get_emp_formatted_name(uid, name, all_uids, day=None):
    """Возвращает имя с цветным значком и меткой 'Под вопросом', если статус MAYBE"""
    clean_name = safe_html(name)
    try:
        idx = all_uids.index(uid) % len(COLOR_BADGES)
        badge = COLOR_BADGES[idx]
    except ValueError:
        badge = "👤"
        
    formatted = f"{badge} <code>{clean_name}</code>"
    
    if day:
        prefs = get_user_preferences(uid)
        if prefs.get(day) == 'MAYBE':
            formatted += " ❓"
            
    return formatted

# ==================== 2. БАЗА ДАННЫХ (SQLITE) ====================

def init_db():
    with sqlite3.connect(DB_NAME) as conn:
        cursor = conn.cursor()
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS employees (
                user_id INTEGER PRIMARY KEY,
                name TEXT NOT NULL,
                shifts_count INTEGER DEFAULT 0
            )
        ''')
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS preferences (
                user_id INTEGER,
                day TEXT NOT NULL,
                status TEXT NOT NULL,
                PRIMARY KEY (user_id, day)
            )
        ''')
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS current_schedule (
                day TEXT,
                user_id INTEGER,
                PRIMARY KEY (day, user_id)
            )
        ''')
        conn.commit()

def register_employee(user_id, name):
    with sqlite3.connect(DB_NAME) as conn:
        cursor = conn.cursor()
        cursor.execute('''
            INSERT INTO employees (user_id, name, shifts_count)
            VALUES (?, ?, 0)
            ON CONFLICT(user_id) DO UPDATE SET name=excluded.name
        ''', (user_id, name))
        conn.commit()

def cycle_preference(user_id, day):
    with sqlite3.connect(DB_NAME) as conn:
        cursor = conn.cursor()
        cursor.execute('SELECT status FROM preferences WHERE user_id = ? AND day = ?', (user_id, day))
        row = cursor.fetchone()
        
        if not row:
            cursor.execute('INSERT INTO preferences (user_id, day, status) VALUES (?, ?, ?)', (user_id, day, 'YES'))
        elif row[0] == 'YES':
            cursor.execute('UPDATE preferences SET status = ? WHERE user_id = ? AND day = ?', ('MAYBE', user_id, day))
        else:
            cursor.execute('DELETE FROM preferences WHERE user_id = ? AND day = ?', (user_id, day))
        conn.commit()

def get_user_preferences(user_id):
    with sqlite3.connect(DB_NAME) as conn:
        cursor = conn.cursor()
        cursor.execute('SELECT day, status FROM preferences WHERE user_id = ?', (user_id,))
        return {row[0]: row[1] for row in cursor.fetchall()}

def get_all_employees():
    with sqlite3.connect(DB_NAME) as conn:
        cursor = conn.cursor()
        cursor.execute('SELECT user_id, name, shifts_count FROM employees')
        return {row[0]: {"name": row[1], "shifts_count": row[2]} for row in cursor.fetchall()}

def update_shift_count(user_id, new_count):
    with sqlite3.connect(DB_NAME) as conn:
        cursor = conn.cursor()
        cursor.execute('UPDATE employees SET shifts_count = ? WHERE user_id = ?', (max(0, new_count), user_id))
        conn.commit()

def reset_all_shifts():
    with sqlite3.connect(DB_NAME) as conn:
        cursor = conn.cursor()
        cursor.execute('UPDATE employees SET shifts_count = 0')
        conn.commit()

def save_schedule_to_db(schedule):
    with sqlite3.connect(DB_NAME) as conn:
        cursor = conn.cursor()
        cursor.execute('DELETE FROM current_schedule')
        for day, uids in schedule.items():
            for uid in uids:
                cursor.execute('INSERT INTO current_schedule (day, user_id) VALUES (?, ?)', (day, uid))
        conn.commit()

def get_schedule_from_db():
    with sqlite3.connect(DB_NAME) as conn:
        cursor = conn.cursor()
        cursor.execute('SELECT day, user_id FROM current_schedule')
        schedule = {}
        for day, uid in cursor.fetchall():
            if day not in schedule:
                schedule[day] = []
            schedule[day].append(uid)
        return schedule

def clear_all_preferences():
    with sqlite3.connect(DB_NAME) as conn:
        cursor = conn.cursor()
        cursor.execute('DELETE FROM preferences')
        conn.commit()

# ==================== 3. КЛАВИАТУРЫ ====================

def get_days_keyboard(user_id):
    markup = types.ReplyKeyboardMarkup(resize_keyboard=True, row_width=2)
    prefs = get_user_preferences(user_id)
    week_days = get_next_week_days()
    
    buttons = []
    for day in week_days:
        status = prefs.get(day)
        if status == 'YES':
            label = f"✅ {day}"
        elif status == 'MAYBE':
            label = f"❓ {day}"
        else:
            label = day
        buttons.append(types.KeyboardButton(label))
    
    markup.add(*buttons)
    markup.add(types.KeyboardButton("🟢 Сохранить пожелания"), types.KeyboardButton("👤 Мой статус"))
    markup.add(types.KeyboardButton("📋 Посмотреть график"), types.KeyboardButton("🔄 Запросить замену"))
    
    if user_id in ADMIN_IDS:
        markup.add(types.KeyboardButton("📊 Сгенерировать график"), types.KeyboardButton("✏️ Ручная правка"))
        markup.add(types.KeyboardButton("⚙️ Управление сменами"), types.KeyboardButton("📥 Скачать Excel"))
        
    return markup

def get_manage_shifts_markup():
    employees = get_all_employees()
    markup = types.InlineKeyboardMarkup()
    markup.add(types.InlineKeyboardButton("🔄 СБРОСИТЬ ВСЕ СМЕНЫ НА 0", callback_data="reset_all_shifts"))
    markup.add(types.InlineKeyboardButton("🗑 СБРОСИТЬ ПОЖЕЛАНИЯ ВСЕХ", callback_data="clear_all_prefs"))
    
    for uid, data in employees.items():
        btn_text = f"{data['name']} ({data['shifts_count']} смен)"
        markup.add(
            types.InlineKeyboardButton(f"➕ {btn_text}", callback_data=f"add_sh:{uid}"),
            types.InlineKeyboardButton(f"➖ {btn_text}", callback_data=f"sub_sh:{uid}")
        )
    return markup

# ==================== 4. ОБРАБОТЧИКИ ====================

@bot.message_handler(commands=['start'])
def start_message(message):
    user_id = message.from_user.id
    name = message.from_user.first_name or f"Сотрудник #{user_id}"
    register_employee(user_id, name)
    
    role = get_user_role_text(user_id)
    safe_send_message(
        message.chat.id,
        f"Привет, <b>{safe_html(name)}</b>!\nВаш статус: {role}\n\n"
        "📌 <b>Как отмечать дни для работы:</b>\n"
        "• Без отметки = <b>Не могу работать</b> (выходной)\n"
        "• 1 клик = ✅ <b>Могу работать</b>\n"
        "• 2 клика = ❓ <b>Под вопросом</b>\n"
        "• 3 клика = <b>Сбросить отметку</b>\n\n"
        "После выбора нажмите <b>«🟢 Сохранить пожелания»</b>.",
        reply_markup=get_days_keyboard(user_id),
        parse_mode="HTML"
    )

@bot.message_handler(func=lambda msg: msg.text and "Мой статус" in msg.text)
def show_status(message):
    user_id = message.from_user.id
    employees = get_all_employees()
    all_uids = list(employees.keys())
    emp_data = employees.get(user_id, {"name": "Сотрудник", "shifts_count": 0})
    
    formatted_name = get_emp_formatted_name(user_id, emp_data['name'], all_uids)
    role = get_user_role_text(user_id)
    
    text = (
        f"📋 <b>Ваша карточка:</b>\n\n"
        f"• <b>Имя:</b> {formatted_name}\n"
        f"• <b>ID:</b> <code>{user_id}</code>\n"
        f"• <b>Роль:</b> {role}\n"
        f"• <b>Отработано смен:</b> {emp_data['shifts_count']}"
    )
    safe_send_message(message.chat.id, text, parse_mode="HTML")

@bot.message_handler(func=lambda msg: msg.text and "Посмотреть график" in msg.text)
def view_schedule(message):
    try:
        schedule = get_schedule_from_db()
        employees = get_all_employees()
        all_uids = list(employees.keys())
        week_days = get_next_week_days()

        response = f"{DIVIDER_LINE}\n🗓 <b>АКТУАЛЬНЫЙ ГРАФИК</b>\n{DIVIDER_LINE}\n\n"

        for day in week_days:
            uids = schedule.get(day, [])
            needed = 2 if ("Суббота" in day or "Воскресенье" in day) else 1
            
            valid_uids = [uid for uid in uids if uid in employees]
            
            if valid_uids:
                formatted_workers = [
                    get_emp_formatted_name(uid, employees[uid]["name"], all_uids, day=day) 
                    for uid in valid_uids
                ]
                workers_str = ", ".join(formatted_workers)
                if len(valid_uids) < needed:
                    workers_str += f" ⚠️ (нехватка: {len(valid_uids)}/{needed})"
            else:
                workers_str = "🔴 Смена еще не сформирована"
                
            response += f"📌 <b>{safe_html(day)}</b>: {workers_str}\n"

        response += f"\n<i>(По вопросам или проблемам писать в личные)</i>\n{DIVIDER_LINE}"
        safe_send_message(message.chat.id, response, parse_mode="HTML")
    except Exception as e:
        safe_send_message(message.chat.id, f"❌ Ошибка при отображении графика: {safe_html(str(e))}")

@bot.message_handler(func=lambda msg: msg.text and any(d.split(" ")[0] in msg.text for d in ["Понедельник", "Вторник", "Среда", "Четверг", "Пятница", "Суббота", "Воскресенье"]))
def handle_day_selection(message):
    user_id = message.from_user.id
    raw_text = message.text.replace("✅ ", "").replace("❓ ", "")
    
    week_days = get_next_week_days()
    if raw_text in week_days:
        cycle_preference(user_id, raw_text)
        safe_send_message(message.chat.id, f"Статус для «{raw_text}» обновлен.", reply_markup=get_days_keyboard(user_id))

@bot.message_handler(func=lambda msg: msg.text and "Сохранить пожелания" in msg.text)
def save_preferences(message):
    user_id = message.from_user.id
    prefs = get_user_preferences(user_id)
    
    if prefs:
        text = "✅ <b>Ваши рабочие дни сохранены:</b>\n\n"
        for day, status in prefs.items():
            st_text = "✅ Готов работать" if status == 'YES' else "❓ Под вопросом"
            text += f"• {safe_html(day)}: {st_text}\n"
        safe_send_message(message.chat.id, text, parse_mode="HTML")
    else:
        safe_send_message(message.chat.id, "ℹ️ Вы не выбрали ни одного дня. Бот считает, что вы <b>отдыхаете на этой неделе</b>.", parse_mode="HTML")

# --- СИСТЕМА ЗАМЕН ---

@bot.message_handler(func=lambda msg: msg.text and "Запросить замену" in msg.text)
def request_swap_start(message):
    markup = types.ReplyKeyboardMarkup(resize_keyboard=True, row_width=2)
    buttons = [types.KeyboardButton(f"Ищу замену: {day}") for day in get_next_week_days()]
    markup.add(*buttons)
    markup.add(types.KeyboardButton("⬅️ Назад в меню"))
    
    safe_send_message(message.chat.id, "Выберите день, на который ищете подмену:", reply_markup=markup)

@bot.message_handler(func=lambda msg: msg.text and msg.text.startswith("Ищу замену: "))
def process_swap_request(message):
    user_id = message.from_user.id
    employees = get_all_employees()
    all_uids = list(employees.keys())
    
    sender_name = employees.get(user_id, {}).get("name", "Сотрудник")
    target_day = message.text.replace("Ищу замену: ", "")
    week_days = get_next_week_days()
    
    if target_day not in week_days:
        safe_send_message(message.chat.id, "Ошибка выбора дня.")
        return
        
    day_idx = week_days.index(target_day)
    formatted_sender = get_emp_formatted_name(user_id, sender_name, all_uids)
    
    alert_text = (
        f"🚨 <b>ЗАПРОС НА ОБМЕН СМЕНОЙ!</b>\n\n"
        f"Сотрудник {formatted_sender} ищет замену на день:\n📅 <b>{safe_html(target_day)}</b>\n\n"
        f"Если вы готовы выйти вместо коллеги, нажмите кнопку ниже:"
    )
    
    markup = types.InlineKeyboardMarkup()
    markup.add(types.InlineKeyboardButton("🙋‍♂️ Я могу подменить", callback_data=f"sw_resp:{user_id}:{day_idx}"))
    
    # Безопасная рассылка пользователям
    for emp_id in all_uids:
        if emp_id != user_id:
            safe_send_message(emp_id, alert_text, reply_markup=markup, parse_mode="HTML")
                
    safe_send_message(message.chat.id, f"✅ Запрос на замену на <b>{safe_html(target_day)}</b> отправлен коллегам!", reply_markup=get_days_keyboard(user_id), parse_mode="HTML")

@bot.callback_query_handler(func=lambda call: call.data.startswith("sw_resp:"))
def handle_swap_response(call):
    _, sender_id_str, day_idx_str = call.data.split(":")
    sender_id = int(sender_id_str)
    day_idx = int(day_idx_str)
    cover_id = call.from_user.id
    
    if cover_id == sender_id:
        bot.answer_callback_query(call.id, "Вы не можете подменить самого себя!", show_alert=True)
        return

    employees = get_all_employees()
    all_uids = list(employees.keys())
    
    week_days = get_next_week_days()
    day_name = week_days[day_idx]
    
    sender_name = employees.get(sender_id, {}).get("name", "Сотрудник 1")
    cover_name = employees.get(cover_id, {}).get("name", "Сотрудник 2")
    
    fmt_sender = get_emp_formatted_name(sender_id, sender_name, all_uids)
    fmt_cover = get_emp_formatted_name(cover_id, cover_name, all_uids)
    
    bot.answer_callback_query(call.id, "Ваша заявка отправлена администратору!")
    safe_send_message(call.message.chat.id, f"👍 Ваш отклик на подмену ({fmt_sender} на {safe_html(day_name)}) отправлен администратору на утверждение.", parse_mode="HTML")
    
    admin_markup = types.InlineKeyboardMarkup()
    admin_markup.add(
        types.InlineKeyboardButton("✅ Утвердить", callback_data=f"sw_app:{sender_id}:{cover_id}:{day_idx}"),
        types.InlineKeyboardButton("❌ Отклонить", callback_data=f"sw_rej:{sender_id}:{cover_id}:{day_idx}")
    )
    
    admin_text = (
        f"📩 <b>ЗАПРОС НА ПОДТВЕРЖДЕНИЕ ЗАМЕНЫ</b>\n\n"
        f"• День: <b>{safe_html(day_name)}</b>\n"
        f"• Кто ищет замену: {fmt_sender}\n"
        f"• Кто готов подменить: {fmt_cover}\n\n"
        f"Утвердить изменение в графике?"
    )
    
    # Безопасная отправка всем администраторам
    for admin_id in ADMIN_IDS:
        safe_send_message(admin_id, admin_text, reply_markup=admin_markup, parse_mode="HTML")

@bot.callback_query_handler(func=lambda call: call.data.startswith("sw_app:") or call.data.startswith("sw_rej:"))
def handle_admin_swap_decision(call):
    if call.from_user.id not in ADMIN_IDS:
        return
        
    action, sender_id_str, cover_id_str, day_idx_str = call.data.split(":")
    sender_id = int(sender_id_str)
    cover_id = int(cover_id_str)
    day_idx = int(day_idx_str)
    
    week_days = get_next_week_days()
    day_name = week_days[day_idx]
    
    employees = get_all_employees()
    all_uids = list(employees.keys())
    
    sender_name = employees.get(sender_id, {}).get("name", "Сотрудник 1")
    cover_name = employees.get(cover_id, {}).get("name", "Сотрудник 2")
    
    fmt_sender = get_emp_formatted_name(sender_id, sender_name, all_uids)
    fmt_cover = get_emp_formatted_name(cover_id, cover_name, all_uids)
    
    if action == "sw_app":
        schedule = get_schedule_from_db()
        day_workers = schedule.get(day_name, [])
        
        if sender_id in day_workers:
            day_workers.remove(sender_id)
            update_shift_count(sender_id, employees[sender_id]["shifts_count"] - 1)
            
        if cover_id not in day_workers:
            day_workers.append(cover_id)
            update_shift_count(cover_id, employees[cover_id]["shifts_count"] + 1)
            
        schedule[day_name] = day_workers
        save_schedule_to_db(schedule)
        
        bot.edit_message_text(f"✅ <b>Замена утверждена!</b>\n{fmt_cover} выходит вместо {fmt_sender} на <b>{safe_html(day_name)}</b>.", call.message.chat.id, call.message.message_id, parse_mode="HTML")
        
        safe_send_message(sender_id, f"🎉 Ваша замена на <b>{safe_html(day_name)}</b> утверждена! Вас подменит {fmt_cover}.", parse_mode="HTML")
        safe_send_message(cover_id, f"✅ Вы успешно утверждены на замену {fmt_sender} на <b>{safe_html(day_name)}</b>!", parse_mode="HTML")
    else:
        bot.edit_message_text("❌ <b>Замена отклонена.</b>", call.message.chat.id, call.message.message_id, parse_mode="HTML")
        safe_send_message(cover_id, f"❌ Администратор отклонил замену на <b>{safe_html(day_name)}</b>.", parse_mode="HTML")

@bot.message_handler(func=lambda msg: msg.text and "Назад в меню" in msg.text)
def back_to_menu(message):
    user_id = message.from_user.id
    safe_send_message(message.chat.id, "Главное меню:", reply_markup=get_days_keyboard(user_id))

# --- УПРАВЛЕНИЕ СМЕНАМИ И EXCEL ---

@bot.message_handler(func=lambda msg: msg.text and "Скачать Excel" in msg.text)
def export_excel(message):
    user_id = message.from_user.id
    if user_id not in ADMIN_IDS:
        safe_send_message(message.chat.id, "⛔ Отказано в доступе.")
        return

    try:
        wb = openpyxl.Workbook()
        
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

        ws1 = wb.active
        ws1.title = "Расписание"
        ws1.append(["День недели / Дата", "Сотрудники на смене"])
        
        for col in range(1, 3):
            cell = ws1.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        week_days = get_next_week_days()
        schedule = get_schedule_from_db()
        employees = get_all_employees()
        
        for day in week_days:
            uids = schedule.get(day, [])
            names = []
            for u in uids:
                if u in employees:
                    pref = get_user_preferences(u).get(day)
                    status_mark = " (под вопросом)" if pref == 'MAYBE' else ""
                    names.append(f"{employees[u]['name']}{status_mark}")
                    
            names_str = ", ".join(names) if names else "Нет сотрудников"
            ws1.append([day, names_str])
            
        for row in ws1.iter_rows(min_row=1, max_row=len(week_days)+1, min_col=1, max_col=2):
            for cell in row:
                cell.border = border

        ws1.column_dimensions['A'].width = 25
        ws1.column_dimensions['B'].width = 45

        ws2 = wb.create_sheet(title="Статистика смен")
        ws2.append(["ID Сотрудника", "Имя сотрудника", "Количество смен"])
        
        for col in range(1, 4):
            cell = ws2.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        row_idx = 2
        for uid, data in employees.items():
            ws2.append([uid, data['name'], data['shifts_count']])
            row_idx += 1
            
        for row in ws2.iter_rows(min_row=1, max_row=row_idx-1, min_col=1, max_col=3):
            for cell in row:
                cell.border = border

        ws2.column_dimensions['A'].width = 18
        ws2.column_dimensions['B'].width = 28
        ws2.column_dimensions['C'].width = 20

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        
        bot.send_document(
            message.chat.id,
            document=('Schedule_Report.xlsx', stream),
            caption="📊 <b>Актуальный отчет по расписанию и сменам</b>",
            parse_mode="HTML"
        )
    except Exception as e:
        safe_send_message(message.chat.id, f"❌ Ошибка генерации Excel: {safe_html(str(e))}")

@bot.message_handler(func=lambda msg: msg.text and "Управление сменами" in msg.text)
def manage_shifts_menu(message):
    user_id = message.from_user.id
    if user_id not in ADMIN_IDS:
        safe_send_message(message.chat.id, "⛔ Отказано в доступе.")
        return

    safe_send_message(
        message.chat.id, 
        "⚙️ <b>Управление количеством смен:</b>\n\nВы можете ручным нажатием добавлять/убавлять смены или полностью сбросить счетчик.", 
        reply_markup=get_manage_shifts_markup(), 
        parse_mode="HTML"
    )

@bot.callback_query_handler(func=lambda call: call.data == "reset_all_shifts")
def callback_reset_all_shifts(call):
    if call.from_user.id not in ADMIN_IDS:
        bot.answer_callback_query(call.id, "⛔ Отказано в доступе.", show_alert=True)
        return
    
    reset_all_shifts()
    bot.answer_callback_query(call.id, "Все смены успешно сброшены до 0!")
    bot.edit_message_text(
        "✅ <b>Все смены сброшены на 0!</b>", 
        call.message.chat.id, 
        call.message.message_id,
        parse_mode="HTML"
    )

@bot.callback_query_handler(func=lambda call: call.data == "clear_all_prefs")
def callback_clear_all_prefs(call):
    if call.from_user.id not in ADMIN_IDS:
        bot.answer_callback_query(call.id, "⛔ Отказано в доступе.", show_alert=True)
        return
    
    clear_all_preferences()
    bot.answer_callback_query(call.id, "Пожелания всех пользователей очищены!")
    bot.edit_message_text(
        "🗑 <b>Все пожелания очищены!</b> Теперь пользователи могут заново проставлять галочки.", 
        call.message.chat.id, 
        call.message.message_id,
        parse_mode="HTML"
    )

@bot.callback_query_handler(func=lambda call: call.data.startswith("add_sh:") or call.data.startswith("sub_sh:"))
def callback_change_shift_count(call):
    if call.from_user.id not in ADMIN_IDS:
        bot.answer_callback_query(call.id, "⛔ Отказано в доступе.", show_alert=True)
        return
        
    action, uid_str = call.data.split(":")
    uid = int(uid_str)
    employees = get_all_employees()
    
    if uid in employees:
        current = employees[uid]["shifts_count"]
        new_val = current + 1 if action == "add_sh" else current - 1
        update_shift_count(uid, new_val)
        
        bot.answer_callback_query(call.id, "Смены обновлены")
        
        try:
            bot.edit_message_reply_markup(
                chat_id=call.message.chat.id, 
                message_id=call.message.message_id, 
                reply_markup=get_manage_shifts_markup()
            )
        except Exception:
            pass

# --- ГЕНЕРАЦИЯ ГРАФИКА И АВТО-РАССЫЛКА ---

@bot.message_handler(func=lambda msg: msg.text and "Сгенерировать график" in msg.text)
def create_schedule(message):
    user_id = message.from_user.id
    if user_id not in ADMIN_IDS:
        safe_send_message(message.chat.id, "⛔ Отказано в доступе.")
        return

    try:
        employees = get_all_employees()
        if not employees:
            safe_send_message(
                message.chat.id, 
                "⚠️ <b>База сотрудников пуста!</b>\n\nПопросите сотрудников зарегистрироваться, перейдя в бота и нажав <b>/start</b>.", 
                parse_mode="HTML"
            )
            return

        all_uids = list(employees.keys())
        week_days = get_next_week_days()
        schedule_db_data = {}

        for day in week_days:
            needed = 2 if ("Суббота" in day or "Воскресенье" in day) else 1
            assigned = []

            available_ready = []
            available_maybe = []

            for uid in all_uids:
                prefs = get_user_preferences(uid)
                status = prefs.get(day)
                if status == 'YES':
                    available_ready.append(uid)
                elif status == 'MAYBE':
                    available_maybe.append(uid)

            def priority_score(uid):
                return employees[uid]["shifts_count"] + random.uniform(0.0, 1.5)

            available_ready.sort(key=priority_score)
            newly_assigned = available_ready[:needed]

            if len(newly_assigned) < needed:
                needed_still = needed - len(newly_assigned)
                available_maybe.sort(key=priority_score)
                newly_assigned.extend(available_maybe[:needed_still])

            assigned.extend(newly_assigned)

            for uid in newly_assigned:
                employees[uid]["shifts_count"] += 1

            schedule_db_data[day] = assigned

        for uid, data in employees.items():
            update_shift_count(uid, data["shifts_count"])

        save_schedule_to_db(schedule_db_data)

        response = f"{DIVIDER_LINE}\n🗓 <b>РАСПИСАНИЕ НА НЕДЕЛЮ</b>\n{DIVIDER_LINE}\n\n"
        
        for day in week_days:
            uids = schedule_db_data.get(day, [])
            needed = 2 if ("Суббота" in day or "Воскресенье" in day) else 1
            if uids:
                formatted_workers = [get_emp_formatted_name(uid, employees[uid]["name"], all_uids, day=day) for uid in uids if uid in employees]
                workers_str = ", ".join(formatted_workers)
                if len(uids) < needed:
                    workers_str += f" ⚠️ (нехватка: {len(uids)}/{needed})"
            else:
                workers_str = "🔴 Нет доступных сотрудников!"
            response += f"📌 <b>{safe_html(day)}</b>: {workers_str}\n"

        response += f"\n{DIVIDER_LINE}\n📊 <b>ИТОГОВАЯ СТАТИСТИКА СМЕН:</b>\n"
        for uid in all_uids:
            data = employees[uid]
            fname = get_emp_formatted_name(uid, data['name'], all_uids)
            response += f"• {fname}: <code>{data['shifts_count']}</code> смен\n"
            
        response += f"\n<i>(Проверяйте график)</i>\n{DIVIDER_LINE}"

        safe_send_message(message.chat.id, response, parse_mode="HTML")

        # Безопасная массовая рассылка сотрудникам
        success_count = 0
        for emp_id in all_uids:
            if emp_id != user_id:
                res = safe_send_message(emp_id, f"🔔 <b>Опубликовано новое расписание!</b>\n\n{response}", parse_mode="HTML")
                if res:
                    success_count += 1
                    
        safe_send_message(message.chat.id, f"📢 График успешно разослан всем сотрудникам (доставлено: {success_count}).", parse_mode="HTML")

    except Exception as e:
        safe_send_message(message.chat.id, f"❌ <b>Ошибка при генерации графика:</b>\n<code>{safe_html(str(e))}</code>", parse_mode="HTML")

# --- РУЧНАЯ ПРАВКА ГРАФИКА ---

@bot.message_handler(func=lambda msg: msg.text and "Ручная правка" in msg.text)
def edit_schedule_start(message):
    user_id = message.from_user.id
    if user_id not in ADMIN_IDS:
        safe_send_message(message.chat.id, "⛔ Отказано в доступе.")
        return

    week_days = get_next_week_days()
    markup = types.InlineKeyboardMarkup()
    for idx, day in enumerate(week_days):
        markup.add(types.InlineKeyboardButton(f"📅 {day}", callback_data=f"ed_day:{idx}"))
        
    safe_send_message(message.chat.id, "Выберите день для изменения состава смены:", reply_markup=markup)

@bot.callback_query_handler(func=lambda call: call.data.startswith("ed_day:"))
def edit_day_callback(call):
    day_idx = int(call.data.split(":")[1])
    week_days = get_next_week_days()
    day_name = week_days[day_idx]
    
    schedule = get_schedule_from_db()
    assigned_uids = schedule.get(day_name, [])
    employees = get_all_employees()
    all_uids = list(employees.keys())
    
    markup = types.InlineKeyboardMarkup()
    for uid, data in employees.items():
        is_assigned = uid in assigned_uids
        btn_text = f"➖ {data['name']}" if is_assigned else f"➕ {data['name']}"
        markup.add(types.InlineKeyboardButton(btn_text, callback_data=f"t_emp:{day_idx}:{uid}"))
        
    assigned_names = ", ".join([get_emp_formatted_name(uid, employees[uid]["name"], all_uids, day=day_name) for uid in assigned_uids if uid in employees]) if assigned_uids else "Никого"
    text = f"⚙ <b>Редактирование {safe_html(day_name)}</b>\n\nСейчас выходят: {assigned_names}\n\nНажимайте на имена, чтобы добавить или удалить человека из смены:"
    
    bot.edit_message_text(text, call.message.chat.id, call.message.message_id, reply_markup=markup, parse_mode="HTML")

@bot.callback_query_handler(func=lambda call: call.data.startswith("t_emp:"))
def toggle_employee_callback(call):
    _, day_idx_str, uid_str = call.data.split(":")
    day_idx = int(day_idx_str)
    uid = int(uid_str)
    
    week_days = get_next_week_days()
    day_name = week_days[day_idx]
    
    schedule = get_schedule_from_db()
    assigned_uids = schedule.get(day_name, [])
    employees = get_all_employees()
    
    if uid in assigned_uids:
        assigned_uids.remove(uid)
        employees[uid]["shifts_count"] = max(0, employees[uid]["shifts_count"] - 1)
    else:
        assigned_uids.append(uid)
        employees[uid]["shifts_count"] += 1
        
    update_shift_count(uid, employees[uid]["shifts_count"])
    schedule[day_name] = assigned_uids
    save_schedule_to_db(schedule)
    
    edit_day_callback(call)

# ==================== 5. ЗАПУСК И АВТО-ПЕРЕЗАПУСК ====================

if __name__ == '__main__':
    init_db()
    print("Бот успешно запущен и защищен от ошибок...")
    
    while True:
        try:
            bot.polling(none_stop=True, interval=0, timeout=20)
        except Exception as e:
            print(f"Сбой подключения: {e}. Перезапуск через 5 секунд...")
            time.sleep(5)